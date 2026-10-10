"""mk-ppt 스킬용 python-pptx 헬퍼.

공식 API에 없는 기능(슬라이드 삭제/복제/이동)과 반복되는 작업(표, 차트, 구조 덤프)을 제공한다.
"""
import copy
import sys

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Pt

__all__ = [
    "delete_slide", "move_slide", "duplicate_slide",
    "add_table", "add_chart", "dump_presentation", "extract_text",
]

CHART_KINDS = {
    "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
    "bar": XL_CHART_TYPE.BAR_CLUSTERED,
    "line": XL_CHART_TYPE.LINE_MARKERS,
    "pie": XL_CHART_TYPE.PIE,
    "doughnut": XL_CHART_TYPE.DOUGHNUT,
    "area": XL_CHART_TYPE.AREA,
    "radar": XL_CHART_TYPE.RADAR,
}


def delete_slide(prs, index):
    """index번째 슬라이드를 삭제한다. 관계까지 끊어 고아 파트가 남지 않게 한다."""
    sld_id_lst = prs.slides._sldIdLst
    sld_id = list(sld_id_lst)[index]
    prs.part.drop_rel(sld_id.rId)
    sld_id_lst.remove(sld_id)


def move_slide(prs, old_index, new_index):
    """슬라이드 순서를 바꾼다."""
    sld_id_lst = prs.slides._sldIdLst
    items = list(sld_id_lst)
    el = items[old_index]
    sld_id_lst.remove(el)
    sld_id_lst.insert(new_index, el)


def duplicate_slide(prs, index):
    """index번째 슬라이드를 맨 뒤에 복제하고 새 슬라이드를 반환한다.

    도형은 XML째 복사한다. 그림/차트/미디어처럼 외부 관계를 가진 도형은
    관계를 함께 복사하지 않으므로 복제본에서 깨질 수 있다(텍스트/도형 위주 슬라이드에 적합).
    """
    src = prs.slides[index]
    dst = prs.slides.add_slide(src.slide_layout)
    for shp in list(dst.shapes):
        shp._element.getparent().remove(shp._element)
    for shp in src.shapes:
        dst.shapes._spTree.insert_element_before(copy.deepcopy(shp._element), "p:extLst")
    if src.has_notes_slide:
        dst.notes_slide.notes_text_frame.text = src.notes_slide.notes_text_frame.text
    return dst


def add_table(slide, rows, left, top, width, height, header=True,
              font_size=14, header_fill=RGBColor(0x1F, 0x3A, 0x5F)):
    """2차원 리스트로 표를 만든다. header=True면 첫 행을 헤더 색으로 칠한다."""
    n_rows, n_cols = len(rows), max(len(r) for r in rows)
    table = slide.shapes.add_table(n_rows, n_cols, left, top, width, height).table
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            cell = table.cell(r, c)
            cell.text = str(value)
            for p in cell.text_frame.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(font_size)
                    if header and r == 0:
                        run.font.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)
            if header and r == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = header_fill
    return table


def add_chart(slide, kind, categories, series, left, top, width, height,
              title=None, data_labels=False):
    """범주형 차트를 추가한다.

    kind: column/bar/line/pie/doughnut/area/radar
    series: {"계열명": [값, ...]}
    """
    data = CategoryChartData()
    data.categories = list(categories)
    for name, values in series.items():
        data.add_series(name, list(values))
    graphic = slide.shapes.add_chart(CHART_KINDS[kind], left, top, width, height, data)
    chart = graphic.chart
    if title:
        chart.has_title = True
        chart.chart_title.text_frame.text = title
    chart.has_legend = len(series) > 1 or kind in ("pie", "doughnut")
    if data_labels:
        chart.plots[0].has_data_labels = True
    return chart


def extract_text(path):
    """모든 슬라이드의 텍스트를 [[줄, ...], ...] 형태로 반환한다."""
    prs = Presentation(path)
    result = []
    for slide in prs.slides:
        lines = []
        for shp in slide.shapes:
            if shp.has_text_frame:
                lines += [p.text for p in shp.text_frame.paragraphs if p.text.strip()]
            if getattr(shp, "has_table", False) and shp.has_table:
                for row in shp.table.rows:
                    lines.append(" | ".join(c.text for c in row.cells))
        result.append(lines)
    return result


def dump_presentation(path, out=sys.stdout):
    """슬라이드별 레이아웃, 도형, 위치, 텍스트, 노트를 출력한다. 생성 결과 검증용."""
    prs = Presentation(path)
    W, H = prs.slide_width, prs.slide_height
    print(f"{path}: {len(prs.slides)}장, {W / 914400:.2f} x {H / 914400:.2f} in", file=out)
    for i, slide in enumerate(prs.slides):
        print(f"\n[슬라이드 {i}] 레이아웃={slide.slide_layout.name}", file=out)
        for shp in slide.shapes:
            pos = ""
            if shp.left is not None:
                pos = f"({shp.left / 914400:.1f},{shp.top / 914400:.1f},{shp.width / 914400:.1f}x{shp.height / 914400:.1f})"
                if shp.left < 0 or shp.top < 0 or shp.left + shp.width > W or shp.top + shp.height > H:
                    pos += " ⚠화면 밖"
            text = shp.text_frame.text.replace("\n", " / ")[:60] if shp.has_text_frame else ""
            empty = " ⚠빈 플레이스홀더" if shp.is_placeholder and shp.has_text_frame and not text.strip() else ""
            print(f"  - {shp.shape_type} {shp.name} {pos} {text}{empty}", file=out)
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            print(f"  노트: {slide.notes_slide.notes_text_frame.text[:60]}", file=out)


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")  # Windows 콘솔(cp949)에서 한글/기호 깨짐 방지
    if len(sys.argv) != 2:
        sys.exit("사용법: python pptx_helpers.py <파일.pptx>")
    dump_presentation(sys.argv[1])
