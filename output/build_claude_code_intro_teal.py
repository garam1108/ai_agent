"""Claude Code 소개 PPT (8장) - Clean Teal Layout (references/ppt/slide-design/system.md)."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

TEAL = RGBColor(0x00, 0x87, 0x7A)
DTEAL = RGBColor(0x00, 0x6B, 0x63)
LTEAL = RGBColor(0x4D, 0xB6, 0xAC)
PTEAL = RGBColor(0xE0, 0xF2, 0xF1)
YELLOW = RGBColor(0xF5, 0xB7, 0x00)
DARK = RGBColor(0x33, 0x33, 0x33)
GRAY = RGBColor(0x80, 0x80, 0x80)
LGRAY = RGBColor(0xF2, 0xF2, 0xF2)
WHITE = RGBColor(255, 255, 255)
FONT = "Noto Sans KR"

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
W, H = 13.333, 7.5
TOTAL = 8


def shape(s, kind, x, y, w, h, fill):
    sh = s.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = fill
    sh.line.fill.background()
    sh.shadow.inherit = False
    return sh


def rect(s, x, y, w, h, fill):
    return shape(s, MSO_SHAPE.RECTANGLE, x, y, w, h, fill)


def circle(s, x, y, d, fill):
    return shape(s, MSO_SHAPE.OVAL, x, y, d, d, fill)


def text(s, x, y, w, h, t, size=14, color=DARK, bold=False,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    for i, line in enumerate(t if isinstance(t, list) else [t]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run(); r.text = line
        r.font.size = Pt(size); r.font.bold = bold
        r.font.color.rgb = color; r.font.name = FONT
    return tb


def line(s, x1, y1, x2, y2, color=GRAY, width=0.75):
    ln = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    ln.line.color.rgb = color; ln.line.width = Pt(width)
    return ln


def content_slide(title, subtitle, n):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    text(s, 0.6, 0.35, 9.5, 0.75, title, 30, DARK, True, anchor=MSO_ANCHOR.MIDDLE)
    text(s, 0.6, 1.1, 9.5, 0.4, subtitle, 13, GRAY)
    line(s, 0.65, 1.6, W - 0.65, 1.6)
    # logo (diamond + MY LOGO)
    shape(s, MSO_SHAPE.DIAMOND, 11.0, 0.55, 0.3, 0.3, TEAL)
    text(s, 11.35, 0.5, 1.4, 0.4, "MY LOGO", 11, TEAL, True, anchor=MSO_ANCHOR.MIDDLE)
    # footer
    rect(s, 0, H - 0.15, W, 0.15, TEAL)
    text(s, 11.7, 6.85, 1.0, 0.35, f"{n} / {TOTAL}", 10, GRAY, align=PP_ALIGN.RIGHT)
    text(s, 0.6, 6.85, 4, 0.35, "claude.com/claude-code", 10, GRAY)
    return s


def num_circle(s, x, y, d, label, size=18, fill=TEAL):
    c = circle(s, x, y, d, fill)
    tf = c.text_frame; tf.margin_left = tf.margin_right = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = label
    r.font.size = Pt(size); r.font.bold = True; r.font.color.rgb = WHITE; r.font.name = FONT
    return c


# ---------- 1. Cover ----------
s = prs.slides.add_slide(prs.slide_layouts[6])
circle(s, 7.0, 3.4, 5.6, PTEAL)
circle(s, 8.0, -1.8, 7.2, LTEAL)
circle(s, 8.8, -1.4, 7.4, TEAL)
circle(s, 10.6, 4.4, 3.4, DTEAL)
shape(s, MSO_SHAPE.DIAMOND, 0.8, 0.6, 0.32, 0.32, TEAL)
text(s, 1.17, 0.55, 2, 0.42, "MY LOGO", 12, TEAL, True, anchor=MSO_ANCHOR.MIDDLE)
text(s, 0.8, 2.3, 6.4, 1.0, "Claude Code 소개", 40, DARK, True, anchor=MSO_ANCHOR.MIDDLE)
rect(s, 0.9, 3.4, 1.2, 0.07, YELLOW)
text(s, 0.8, 3.65, 6.4, 0.9, ["터미널에서 일하는", "AI 코딩 에이전트"], 20, DTEAL, True)
text(s, 0.8, 5.9, 5, 0.4, "claude.com/claude-code", 12, GRAY)

# ---------- 2. What is Claude Code ----------
s = content_slide("Claude Code란?", "코드베이스를 이해하고 직접 작업하는 에이전트", 2)
rect(s, 0.6, 1.95, 4.6, 4.6, PTEAL)
text(s, 0.85, 2.2, 4.1, 1.2, ["터미널 속", "AI 코딩 에이전트"], 22, DTEAL, True)
text(s, 0.85, 3.6, 4.1, 2.6,
     ["자연어로 요청하면 코드를 읽고,", "파일을 고치고, 명령을 실행합니다.", "", "설명만 하는 챗봇이 아니라", "직접 일하는 동료에 가깝습니다."], 14, DARK)
items = [("코드 이해", "프로젝트 전체 구조를 파악합니다."),
         ("파일 수정", "여러 파일을 한 번에 편집합니다."),
         ("명령 실행", "테스트와 빌드를 직접 돌립니다."),
         ("대화형 작업", "결과를 보며 계속 다듬습니다.")]
for i, (t, d) in enumerate(items):
    cx = 5.6 + (i % 2) * 3.65
    cy = 1.95 + (i // 2) * 2.4
    rect(s, cx, cy, 3.45, 2.2, LGRAY)
    num_circle(s, cx + 0.25, cy + 0.25, 0.7, str(i + 1))
    text(s, cx + 0.2, cy + 1.05, 3.05, 0.4, t, 18, DARK, True)
    text(s, cx + 0.2, cy + 1.5, 3.05, 0.6, d, 12, DARK)

# ---------- 3. Hub: core abilities ----------
s = content_slide("핵심 기능", "하나의 에이전트가 개발 흐름 전체를 다룹니다", 3)
hx, hy = W / 2, 4.3
cards = [(0.6, 1.95, "코드 탐색", "파일과 코드를 검색하고\n읽어서 맥락을 파악합니다.", (5.0, 3.35)),
         (8.33, 1.95, "파일 편집", "요청에 맞게 코드를\n정확히 수정합니다.", (8.33, 3.35)),
         (0.6, 5.0, "명령 실행", "테스트, 빌드 등 터미널\n명령을 실행합니다.", (5.0, 5.0)),
         (8.33, 5.0, "Git 작업", "변경 내용을 확인하고\n커밋 작업을 돕습니다.", (8.33, 5.0))]
for x, y, t, d, (ex, ey) in cards:
    line(s, hx, hy, ex, ey, LTEAL, 1)
for i, (x, y, t, d, _) in enumerate(cards):
    rect(s, x, y, 4.4, 1.4, LGRAY)
    num_circle(s, x + 0.2, y + 0.3, 0.8, str(i + 1))
    text(s, x + 1.2, y + 0.12, 3.1, 0.4, t, 18, DARK, True)
    text(s, x + 1.2, y + 0.55, 3.1, 0.75, d.split("\n"), 12, DARK)
c = circle(s, hx - 1.0, hy - 1.0, 2.0, TEAL)
circle(s, hx - 0.8, hy - 0.8, 1.6, DTEAL)
text(s, hx - 0.8, hy - 0.5, 1.6, 1.0, ["Claude", "Code"], 16, WHITE, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)

# ---------- 4. Steps ----------
s = content_slide("작업 흐름", "요청부터 검토까지 네 단계로 진행됩니다", 4)
steps = [("요청", "하고 싶은 일을\n자연어로 말합니다."),
         ("탐색", "관련 파일을 찾아\n코드를 읽습니다."),
         ("수정", "계획에 따라\n파일을 편집합니다."),
         ("검토", "변경 내용을 확인하고\n승인 또는 수정합니다.")]
sw, gap = 2.81, 0.3
line(s, 0.9, 2.6, W - 0.9, 2.6, LTEAL, 1)
for i, (t, d) in enumerate(steps):
    x = 0.6 + i * (sw + gap)
    rect(s, x, 2.0, sw, 3.9, LGRAY)
    num_circle(s, x + sw / 2 - 0.45, 2.15, 0.9, f"{i + 1}", 22, YELLOW if i == 3 else TEAL)
    text(s, x, 3.3, sw, 0.35, f"STEP 0{i + 1}", 12, GRAY, False, PP_ALIGN.CENTER)
    text(s, x, 3.75, sw, 0.5, t, 20, DTEAL, True, PP_ALIGN.CENTER)
    text(s, x + 0.2, 4.5, sw - 0.4, 1.2, d.split("\n"), 14, DARK, False, PP_ALIGN.CENTER)

# ---------- 5. Tools & permissions ----------
s = content_slide("도구와 권한", "승인 단계로 안전하게 작업합니다", 5)
rows = [("읽기 도구", "Read, Grep, Glob으로 파일과 코드를 찾아 읽습니다."),
        ("편집 도구", "Edit, Write로 파일을 수정하며 실행 전 승인을 받습니다."),
        ("명령 실행", "터미널 명령을 실행하고 허용 규칙으로 승인을 줄입니다."),
        ("권한 모드", "승인 방식을 골라 안전성과 속도를 조절합니다.")]
for i, (lab, desc) in enumerate(rows):
    y = 1.95 + i * 1.2
    rect(s, 0.6, y, 3.0, 1.0, TEAL)
    text(s, 0.6, y, 3.0, 1.0, lab, 18, WHITE, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    rect(s, 3.6, y, 9.13, 1.0, LGRAY)
    text(s, 3.85, y, 8.7, 1.0, desc, 14, DARK, False, anchor=MSO_ANCHOR.MIDDLE)

# ---------- 6. Extensions ----------
s = content_slide("확장 기능", "프로젝트에 맞게 Claude Code를 넓힙니다", 6)
exts = [("Skills", "반복 작업의 절차를\n패키지로 재사용합니다."),
        ("Rules", "CLAUDE.md로\n프로젝트 지침을\n전달합니다."),
        ("Hooks", "특정 시점에 명령을\n자동으로 실행합니다."),
        ("MCP", "외부 도구와 서비스를\n연결합니다."),
        ("Subagents", "작업을 별도 에이전트에\n위임합니다.")]
cw, cg = 2.226, 0.25
for i, (t, d) in enumerate(exts):
    x = 0.6 + i * (cw + cg)
    rect(s, x, 2.0, cw, 3.9, PTEAL)
    num_circle(s, x + cw / 2 - 0.4, 2.3, 0.8, str(i + 1))
    text(s, x, 3.35, cw, 0.45, t, 18, DTEAL, True, PP_ALIGN.CENTER)
    text(s, x + 0.1, 4.0, cw - 0.2, 1.8, d.split("\n"), 14, DARK, False, PP_ALIGN.CENTER)

# ---------- 7. Environments ----------
s = content_slide("사용 환경", "익숙한 곳에서 같은 Claude Code를 씁니다", 7)
envs = [("터미널", "CLI에서 바로\n실행합니다."),
        ("데스크톱 앱", "Mac과 Windows용\n앱을 제공합니다."),
        ("웹 앱", "claude.ai/code에서\n브라우저로 사용합니다."),
        ("IDE 확장", "VS Code, JetBrains\n안에서 사용합니다.")]
for i, (t, d) in enumerate(envs):
    x = 0.6 + i * (sw + gap)
    rect(s, x, 2.0, sw, 3.9, LGRAY)
    text(s, x + 0.15, 2.2, sw - 0.3, 1.1, f"0{i + 1}", 48, LTEAL, True)
    if i == 0:
        circle(s, x + sw - 0.55, 2.45, 0.25, YELLOW)
    rect(s, x + 0.2, 3.45, 0.6, 0.05, TEAL)
    text(s, x + 0.15, 3.7, sw - 0.3, 0.5, t, 20, DTEAL, True)
    text(s, x + 0.15, 4.4, sw - 0.3, 1.2, d.split("\n"), 14, DARK)

# ---------- 8. Closing ----------
s = prs.slides.add_slide(prs.slide_layouts[6])
bg = rect(s, 0, 0, W, H, TEAL)
band = shape(s, MSO_SHAPE.WAVE, -0.5, 5.3, W + 1.0, 2.6, LTEAL)
band2 = shape(s, MSO_SHAPE.WAVE, -0.5, 6.0, W + 1.0, 2.0, DTEAL)
text(s, 1.5, 2.1, W - 3.0, 1.3, "Thank You", 54, WHITE, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
rect(s, W / 2 - 0.6, 3.6, 1.2, 0.07, YELLOW)
text(s, 1.5, 3.9, W - 3.0, 0.6, "Claude Code 소개", 20, WHITE, False, PP_ALIGN.CENTER)
text(s, 1.5, 4.5, W - 3.0, 0.4, "claude.com/claude-code", 12, PTEAL, False, PP_ALIGN.CENTER)

prs.save("output/claude_code_intro_teal.pptx")
print("saved")
