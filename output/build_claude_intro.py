import sys
sys.path.insert(0, ".claude/skills/mk-ppt/scripts")
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

ORANGE = RGBColor(0xE8, 0x6A, 0x10)
ORANGE_LIGHT = RGBColor(0xFD, 0xEB, 0xDA)
DARK = RGBColor(0x2B, 0x2B, 0x2B)
GRAY = RGBColor(0x6B, 0x6B, 0x6B)
WHITE = RGBColor(255, 255, 255)
FONT = "맑은 고딕"

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
W, H = 13.333, 7.5

def rect(s, x, y, w, h, fill, shape=MSO_SHAPE.RECTANGLE):
    sh = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = fill
    sh.line.fill.background()
    sh.shadow.inherit = False
    return sh

def text(s, x, y, w, h, t, size=18, color=DARK, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    lines = t if isinstance(t, list) else [t]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run(); r.text = line
        r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color; r.font.name = FONT
        if i: p.space_before = Pt(8)
    return tb

def base(title, n, notes):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    rect(s, 0, 0, 0.25, H, ORANGE)
    text(s, 0.8, 0.45, 11.5, 0.9, title, 34, DARK, True, anchor=MSO_ANCHOR.MIDDLE)
    rect(s, 0.85, 1.35, 1.2, 0.06, ORANGE)
    text(s, 12.0, 6.95, 1.0, 0.4, str(n), 12, GRAY, align=PP_ALIGN.RIGHT)
    s.notes_slide.notes_text_frame.text = notes
    return s

def card(s, x, y, w, h, head, body):
    rect(s, x, y, w, h, ORANGE_LIGHT, MSO_SHAPE.ROUNDED_RECTANGLE).adjustments[0] = 0.06
    rect(s, x, y, 0.09, h, ORANGE)
    text(s, x + 0.3, y + 0.2, w - 0.5, 0.6, head, 22, ORANGE, True)
    text(s, x + 0.3, y + 0.9, w - 0.5, h - 1.0, body, 18, DARK)

# 1. 표지
s = prs.slides.add_slide(prs.slide_layouts[6])
rect(s, 0, 0, W, H, ORANGE)
rect(s, 0.9, 3.55, 1.6, 0.08, WHITE)
text(s, 0.9, 1.8, 11, 1.6, "Anthropic & Claude", 54, WHITE, True, anchor=MSO_ANCHOR.BOTTOM)
text(s, 0.9, 3.85, 11, 1.0, "안전하고 유용한 AI를 만드는 회사 소개", 28, WHITE)
text(s, 0.9, 6.3, 11, 0.5, "회사 소개 발표자료", 16, ORANGE_LIGHT)
s.notes_slide.notes_text_frame.text = "Claude를 만든 Anthropic이 어떤 회사인지 소개합니다."

# 2. 회사 개요
s = base("회사 개요: Anthropic", 2, "Anthropic은 2021년 설립된 AI 안전·연구 기업입니다.")
card(s, 0.85, 1.9, 3.8, 4.3, "설립", ["2021년 설립", "본사: 미국 샌프란시스코"])
card(s, 4.8, 1.9, 3.8, 4.3, "형태", ["공익법인(PBC)", "사회적 이익과 사업을 함께 추구"])
card(s, 8.75, 1.9, 3.8, 4.3, "미션", ["AI가 인류에게 안전하고 유익하도록 만드는 것"])

# 3. 제품
s = base("Claude 제품군", 3, "Claude는 Anthropic의 AI 어시스턴트이며 여러 방식으로 제공됩니다.")
card(s, 0.85, 1.9, 3.8, 4.3, "Claude 앱", ["웹·데스크톱·모바일", "대화, 글쓰기, 분석, 문서 작업"])
card(s, 4.8, 1.9, 3.8, 4.3, "Claude Code", ["터미널·IDE에서 쓰는", "코딩 에이전트"])
card(s, 8.75, 1.9, 3.8, 4.3, "Claude API", ["개발자가 자사 서비스에", "Claude를 통합"])

# 4. 안전
s = base("안전에 대한 접근", 4, "안전 연구는 Anthropic의 핵심 정체성입니다.")
items = [("Constitutional AI", "원칙(헌법)에 기반해 모델이 스스로 행동을 개선하도록 학습"),
         ("책임 있는 확장 정책", "모델 능력이 커질수록 더 엄격한 안전 조치를 적용"),
         ("해석가능성 연구", "모델 내부 동작을 이해해 신뢰성을 높이는 연구")]
for i, (h, b) in enumerate(items):
    y = 1.9 + i * 1.5
    c = rect(s, 0.85, y, 0.9, 0.9, ORANGE, MSO_SHAPE.OVAL)
    c.text_frame.text = str(i + 1)
    r = c.text_frame.paragraphs[0].runs[0]
    r.font.size = Pt(28); r.font.bold = True; r.font.color.rgb = WHITE; r.font.name = FONT
    c.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    text(s, 2.0, y - 0.05, 10.5, 0.5, h, 24, ORANGE, True)
    text(s, 2.0, y + 0.5, 10.5, 0.5, b, 18, DARK)

# 5. 마무리
s = base("정리", 5, "핵심 메시지를 요약하며 마무리합니다.")
rect(s, 0.85, 1.9, 11.7, 1.7, ORANGE, MSO_SHAPE.ROUNDED_RECTANGLE).adjustments[0] = 0.1
text(s, 1.2, 1.9, 11.0, 1.7, "안전을 최우선으로, 유용한 AI를 만든다", 32, WHITE, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
text(s, 0.85, 4.0, 11.7, 2.4, [
    "•  Anthropic: 2021년 설립된 AI 안전 중심의 공익법인",
    "•  Claude: 앱, Claude Code, API로 제공되는 AI 어시스턴트",
    "•  안전 연구와 제품 개발을 함께 추진", 
    "감사합니다"], 20, DARK)

prs.save("output/claude_company_intro.pptx")
